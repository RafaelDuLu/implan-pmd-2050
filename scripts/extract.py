"""Extrae valores almacenados; nunca recalcula ni modifica el libro original.
Uso: python scripts/extract.py /ruta/al/archivo.xlsx
Requiere openpyxl solamente en la computadora de actualización.
"""
import sys, json, csv, hashlib
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parents[1]
path = Path(sys.argv[1])
w = openpyxl.load_workbook(path, data_only=True)
formats = openpyxl.load_workbook(path, data_only=False)
obs = {}
for row in w['Observaciones'].iter_rows():
    v = [c.value for c in row]
    if isinstance(v[1], str) and v[1][0:1].isdigit():
        obs[v[1]] = dict(zip(['nombre','notaNombre','notaBase','notaMeta','fuente','enlace'],v[2:8]))
axes, records = [], []
s = w['TC Indicadores por eje']
for row in s.iter_rows():
    v = [c.value for c in row]
    if isinstance(v[1], str) and not v[2]:
        axes.append({'id': len(axes)+1, 'name': v[1].strip()})
    if not isinstance(v[2], str) or not v[2][0:1].isdigit(): continue
    ident = v[2]; o = obs.get(ident, {})
    reasons = []
    for idx, label in [(6,'Año base'),(7,'Valor base'),(8,'Meta 2030'),(10,'Año último valor'),(11,'Último valor')]:
        if not isinstance(v[idx], (int,float)): reasons.append(label+' no disponible')
    if not v[4]: reasons.append('Unidad de medida no declarada')
    if not o.get('fuente'): reasons.append('Fuente no declarada')
    if o.get('notaBase') or o.get('notaMeta') or v[18]: reasons.append('Observaciones sustantivas pendientes de validación')
    rec = dict(id=ident,axis=len(axes),name=' '.join(v[3].split()),unit=v[4],direction=v[5],baseYear=v[6],base=v[7],target2030=v[8],target2050=v[9],year=v[10],value=v[11],nextUpdate=v[12],source=o.get('fuente'),link=o.get('enlace'),notes=[str(x) for x in [o.get('notaNombre'),o.get('notaBase'),o.get('notaMeta'),v[18]] if x],sheet=s.title,row=row[0].row,percent='%' in formats[s.title].cell(row[0].row,8).number_format,eligible=not reasons,reasons=reasons)
    records.append(rec)
leaders=[]
for row in w['TC Indicadores Lider'].iter_rows():
    v=[c.value for c in row]
    if v[2] is not None and str(v[2])[0:1].isdigit(): leaders.append({'id':str(v[2]),'name':v[3],'baseYear':v[6],'base':v[7],'target2050':v[9],'year':v[10],'value':v[11],'sheet':'TC Indicadores Lider','row':row[0].row})
data={'title':'Monitoreo PMD · Visión 2050','version':'Prueba 1 · Septiembre 2026','file':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'axes':axes,'indicators':records,'leaders':leaders,'scope':'León, Guanajuato; ámbito de referencia sujeto a ficha de cada fuente. No contiene desgloses por municipio o sexo.'}
(root/'assets/data.js').write_text('window.PMD_DATA = '+json.dumps(data,ensure_ascii=False,indent=2)+';\n',encoding='utf-8')
(root/'assets/data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
with (root/'docs/inventario.csv').open('w',encoding='utf-8-sig',newline='') as f:
    out=csv.writer(f);out.writerow(['Clave','Eje','Indicador','Incluido prueba','Motivo de reserva','Hoja','Fila'])
    for d in records:out.writerow([d['id'],d['axis'],d['name'],d['eligible'],'; '.join(d['reasons']),d['sheet'],d['row']])
print(json.dumps({'axes':len(axes),'sectoriales':len(records),'lideres':len(leaders),'seleccionados':sum(d['eligible'] for d in records),'por_eje':{a['id']:sum(d['eligible'] and d['axis']==a['id'] for d in records) for a in axes}},ensure_ascii=False))
